from copy import copy
from datetime import datetime
from io import BytesIO
from urllib.parse import quote

from openpyxl import load_workbook


def _parse_date(value):
    raw = (value or "").strip()[:10]
    if not raw:
        return None
    try:
        return datetime.strptime(raw, "%Y-%m-%d").date()
    except ValueError:
        return None


def _dmy(value):
    d = _parse_date(value)
    if not d:
        return ""
    return d.strftime("%d.%m.%Y")


def export_filename(date_from, date_to):
    a, b = _dmy(date_from), _dmy(date_to)
    if a and b:
        return f"Finex_Hisobot ({a}-{b}).xlsx"
    return "Finex_Hisobot.xlsx"


def _snap_row(ws, row, cols):
    out = []
    for c in range(1, cols + 1):
        cell = ws.cell(row, c)
        out.append(
            {
                "font": copy(cell.font),
                "border": copy(cell.border),
                "fill": copy(cell.fill),
                "number_format": cell.number_format,
                "alignment": copy(cell.alignment),
                "protection": copy(cell.protection),
            }
        )
    return out


def _apply_row(ws, row, snap):
    for c, st in enumerate(snap, 1):
        cell = ws.cell(row, c)
        cell.font = copy(st["font"])
        cell.border = copy(st["border"])
        cell.fill = copy(st["fill"])
        cell.number_format = st["number_format"]
        cell.alignment = copy(st["alignment"])
        cell.protection = copy(st["protection"])


def _set(ws, row, col, value, snap=None, snap_col=None):
    cell = ws.cell(row, col, value)
    if snap:
        st = snap[(snap_col or col) - 1]
        cell.font = copy(st["font"])
        cell.border = copy(st["border"])
        cell.fill = copy(st["fill"])
        cell.number_format = st["number_format"]
        cell.alignment = copy(st["alignment"])
        cell.protection = copy(st["protection"])
    return cell


def fill_hisobot_xlsx(template_path, payload, report_no="00001"):
    wb = load_workbook(template_path)
    ws1 = wb.worksheets[0]
    ws2 = wb.worksheets[1]

    date_from = payload.get("date_from") or ""
    date_to = payload.get("date_to") or ""
    shop = payload.get("store_name") or ""
    rows = payload.get("rows") or []
    turn = payload.get("turnover") or []
    opening = float(payload.get("opening_stock") or 0)
    kirim = sum(float(r.get("amount") or 0) for r in rows if r.get("type") == "kirim")
    chiqim = sum(float(r.get("amount") or 0) for r in rows if r.get("type") == "chiqim")
    closing = opening + kirim - chiqim

    data_snap = _snap_row(ws1, 5, 6)
    tot_label_snap = _snap_row(ws1, 20, 1)
    tot_val_snap = _snap_row(ws1, 20, 3)
    prod_snap = _snap_row(ws2, 7, 12)
    jami_snap = _snap_row(ws2, 6, 12)
    sign_snap = _snap_row(ws2, 35, 2)

    ws1["B1"] = str(report_no)
    ws1["B2"] = shop
    d1, d2 = _parse_date(date_from), _parse_date(date_to)
    if d1:
        ws1["F1"] = d1
    if d2:
        ws1["F2"] = d2

    if ws1.max_row >= 5:
        ws1.delete_rows(5, ws1.max_row - 4)

    r = 5
    for item in rows:
        _apply_row(ws1, r, data_snap)
        ws1.cell(r, 1, (item.get("at") or "")[:19].replace("T", " "))
        ws1.cell(r, 2, "Kirim" if item.get("type") == "kirim" else "Chiqim")
        ws1.cell(r, 3, item.get("title") or "")
        ws1.cell(r, 4, float(item.get("qty") or 0))
        money = ws1.cell(r, 5, float(item.get("amount") or 0))
        money.number_format = data_snap[4]["number_format"]
        ws1.cell(r, 6, item.get("ref") or "")
        r += 1
    r += 1
    totals = [
        ("Davr boshidagi qoldiq:", opening),
        ("Hisobot davrida kirim:", round(kirim, 2)),
        ("Hisobot davrida chiqim:", round(chiqim, 2)),
        ("Davr oxiridagi qoldiq:", round(closing, 2)),
    ]
    for label, value in totals:
        _set(ws1, r, 1, label, tot_label_snap, 1)
        val = _set(ws1, r, 3, value, tot_val_snap, 3)
        val.number_format = tot_val_snap[2]["number_format"]
        r += 1

    ws2["C1"] = str(report_no)
    ws2["C2"] = shop
    if d1:
        ws2["H1"] = d1
    if d2:
        ws2["H2"] = d2

    if ws2.max_row >= 7:
        ws2.delete_rows(7, ws2.max_row - 6)

    first_prod = 7
    for i, t in enumerate(turn):
        rr = first_prod + i
        _apply_row(ws2, rr, prod_snap)
        ws2.cell(rr, 1, t.get("n") or (i + 1))
        ws2.cell(rr, 2, t.get("name") or "")
        ws2.cell(rr, 3, t.get("unit") or "")
        ws2.cell(rr, 4, float(t.get("price") or 0))
        ws2.cell(rr, 5, float(t.get("open_qty") or 0))
        ws2.cell(rr, 6, float(t.get("open_sum") or 0))
        ws2.cell(rr, 7, float(t.get("in_qty") or 0))
        ws2.cell(rr, 8, float(t.get("in_sum") or 0))
        ws2.cell(rr, 9, float(t.get("out_qty") or 0))
        ws2.cell(rr, 10, float(t.get("out_sum") or 0))
        ws2.cell(rr, 11, float(t.get("close_qty") or 0))
        ws2.cell(rr, 12, float(t.get("close_sum") or 0))
        for col in (4, 6, 8, 10, 12):
            ws2.cell(rr, col).number_format = prod_snap[col - 1]["number_format"] or '#,##0.00'

    last_prod = first_prod + len(turn) - 1 if turn else 6
    _apply_row(ws2, 6, jami_snap)
    ws2.cell(6, 1, "Jami")
    if turn:
        for col, letter in [(5, "E"), (6, "F"), (7, "G"), (8, "H"), (9, "I"), (10, "J"), (11, "K"), (12, "L")]:
            cell = ws2.cell(6, col, f"=SUM({letter}{first_prod}:{letter}{last_prod})")
            cell.number_format = jami_snap[col - 1]["number_format"]
    else:
        for col in range(5, 13):
            ws2.cell(6, col, 0)
            ws2.cell(6, col).number_format = jami_snap[col - 1]["number_format"]

    sign_row = (last_prod if turn else 6) + 3
    _set(ws2, sign_row, 2, "Do'kon mudiri _________________________________", sign_snap, 2)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def content_disposition(filename):
    ascii_name = "Finex_Hisobot.xlsx"
    return "attachment; filename=\"%s\"; filename*=UTF-8''%s" % (ascii_name, quote(filename))
